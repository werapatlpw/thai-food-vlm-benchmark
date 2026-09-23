import os, re, json, time, base64
from openai import OpenAI
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ===== CONFIG =====
# Set your API key as an environment variable before running:
#   export XAI_API_KEY="xai-..."
XAI_API_KEY  = os.environ["XAI_API_KEY"]
MODEL = "grok-4.3"
IMAGE_ROOT   = "./images"
OUTPUT_FILE  = "./benchmark_results/benchmark_results_grok.xlsx"
RESUME_FILE  = "./benchmark_results/benchmark_progress_grok.json"
DELAY_SEC    = 2
MAX_IMAGES   = 200
# ==================

PROMPT = """ดูรูปอาหารไทยนี้แล้วประเมินคุณค่าทางโภชนาการต่อ 1 จาน

ตอบในรูปแบบนี้เท่านั้น ห้ามเพิ่มข้อความอื่น:

เมนู: [ชื่ออาหารภาษาไทย]
ส่วนประกอบ:
- [วัตถุดิบ]: [ปริมาณ] g
พลังงาน: [ตัวเลข] kcal
โปรตีน: [ตัวเลข] g
ไขมัน: [ตัวเลข] g
คาร์โบไฮเดรต: [ตัวเลข] g"""


def parse_response(text):
    def extract(pattern):
        m = re.search(pattern, text)
        try: return float(m.group(1)) if m else None
        except: return None
    menu = ""
    m = re.search(r"เมนู:\s*(.+)", text)
    if m: menu = m.group(1).strip()
    return {
        "menu_pred":    menu,
        "kcal_pred":    extract(r"พลังงาน:\s*([\d.]+)"),
        "protein_pred": extract(r"โปรตีน:\s*([\d.]+)"),
        "fat_pred":     extract(r"ไขมัน:\s*([\d.]+)"),
        "carb_pred":    extract(r"คาร์โบไฮเดรต:\s*([\d.]+)"),
        "raw": text,
    }


def encode_image(path):
    with open(path, "rb") as f:
        return base64.b64encode(f.read()).decode("utf-8")


def get_all_images(root):
    items = []
    for cat in sorted(os.listdir(root)):
        cat_path = os.path.join(root, cat)
        if not os.path.isdir(cat_path): continue
        for menu in sorted(os.listdir(cat_path)):
            menu_path = os.path.join(cat_path, menu)
            if not os.path.isdir(menu_path): continue
            for img in sorted(f for f in os.listdir(menu_path) if f.endswith(".jpg")):
                items.append((cat, menu, img.replace(".jpg",""), os.path.join(menu_path, img)))
    return items


def run_benchmark():
    client = OpenAI(api_key=XAI_API_KEY, base_url="https://api.x.ai/v1")
    images = get_all_images(IMAGE_ROOT)
    total  = len(images)
    print(f"พบ {total} รูป ({total//5} เมนู x 5 รูป)")

    results   = []
    done_keys = set()
    if os.path.exists(RESUME_FILE):
        with open(RESUME_FILE, "r", encoding="utf-8") as f:
            results = json.load(f)
        done_keys = {(r["category"], str(r["menu_no"]).zfill(2), str(r["img_no"])) for r in results}
        print(f"Resume: {len(done_keys)} รูปที่ทำไปแล้ว")

    new_count = 0
    for i, (cat, menu_no, img_no, path) in enumerate(images, 1):
        if (cat, menu_no.zfill(2), img_no) in done_keys:
            print(f"[{i}/{total}] skip (done)")
            continue
        if new_count >= MAX_IMAGES:
            print(f"\nครบ {MAX_IMAGES} รูปแล้ว -- รันใหม่เพื่อทำต่อ")
            break
        new_count += 1

        print(f"[{i}/{total}] {cat}/{menu_no}/{img_no}.jpg ...", end=" ", flush=True)

        for attempt in range(3):
            try:
                b64  = encode_image(path)
                resp = client.chat.completions.create(
                    model=MODEL,
                    messages=[{
                        "role": "user",
                        "content": [
                            {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b64}"}},
                            {"type": "text", "text": PROMPT},
                        ]
                    }],
                    max_tokens=512,
                )
                text   = resp.choices[0].message.content
                parsed = parse_response(text)
                parsed.update({"category": cat, "menu_no": menu_no, "img_no": img_no})
                results.append(parsed)
                print(f"OK {parsed['kcal_pred']} kcal")
                break
            except Exception as e:
                err = str(e)
                if ("429" in err or "rate" in err.lower()) and attempt < 2:
                    wait = 60 * (attempt + 1)
                    print(f"rate limit รอ {wait}s...", end=" ")
                    time.sleep(wait)
                else:
                    print(f"Error: {e}")
                    results.append({"category": cat, "menu_no": menu_no, "img_no": img_no,
                                    "menu_pred": "", "kcal_pred": None, "protein_pred": None,
                                    "fat_pred": None, "carb_pred": None, "raw": err})
                    break

        with open(RESUME_FILE, "w", encoding="utf-8") as f:
            json.dump(results, f, ensure_ascii=False)

        if new_count < MAX_IMAGES:
            time.sleep(DELAY_SEC)

    save_results(results)
    done = sum(1 for r in results if r.get("kcal_pred") is not None)
    print(f"\nบันทึกแล้ว: {OUTPUT_FILE}  ({done}/{total} สำเร็จ)")


def save_results(results):
    wb  = openpyxl.Workbook()
    ws  = wb.active
    ws.title = "Grok-4 Results"
    headers = ["category","menu_no","img_no","menu_pred","kcal_pred","protein_pred","fat_pred","carb_pred","raw_response"]
    widths  = [14, 9, 7, 28, 12, 13, 10, 13, 60]
    hf    = Font(bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="1DA1F2")  # สีฟ้า xAI/X
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(1, c, h)
        cell.font = hf; cell.fill = hfill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 20
    for i, r in enumerate(results, 2):
        ws.cell(i,1,r.get("category","")); ws.cell(i,2,r.get("menu_no",""))
        ws.cell(i,3,r.get("img_no","")); ws.cell(i,4,r.get("menu_pred",""))
        ws.cell(i,5,r.get("kcal_pred")); ws.cell(i,6,r.get("protein_pred"))
        ws.cell(i,7,r.get("fat_pred")); ws.cell(i,8,r.get("carb_pred"))
        ws.cell(i,9,r.get("raw",""))
    ws.freeze_panes = "A2"
    wb.save(OUTPUT_FILE)


if __name__ == "__main__":
    run_benchmark()
